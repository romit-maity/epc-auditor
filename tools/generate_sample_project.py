"""Generates the sample EPC project as real files (docx contract, PDF specs, PDF drawing sheets, CSV register/schedules)
plus _manifest.json (ground truth of planted deviations). Needs: python-docx reportlab. Run from repo root."""
import csv, json, re, textwrap
from pathlib import Path
import docx
from reportlab.pdfgen import canvas

OUT = Path("backend/corpus/sample_project"); OUT.mkdir(parents=True, exist_ok=True)
g = lambda v: f"{v:g}"
# (group, parameter, unit, spec value, tol, drawing, overrides)  overrides: con=contract value, dwg=drawing value
ITEMS = [
 ("M","casing bore diameter","mm",100.0,0.5,"DWG-114",{"dwg":101.4}),
 ("M","casing wall thickness","mm",9.0,0.1,"DWG-114",{}),
 ("M","impeller diameter","mm",210.0,1.0,"DWG-114",{}),
 ("M","shaft diameter","mm",45.0,0.05,"DWG-115",{"dwg":45.12}),
 ("M","keyway width","mm",14.0,0.05,"DWG-115",{}),
 ("M","bearing housing bore","mm",85.0,0.03,"DWG-115",{}),
 ("M","motor rated power","kW",75,0,"DWG-204",{}),
 ("M","motor baseplate thickness","mm",20.0,0.5,"DWG-204",{"dwg":21.2}),
 ("M","coupling spacer length","mm",140,1,"DWG-204",{}),
 ("M","suction nozzle size","mm",150,0,"DWG-203",{}),
 ("M","discharge nozzle size","mm",100,0,"DWG-203",{}),
 ("M","nozzle projection","mm",300,2,"DWG-203",{}),
 ("P","design pressure","bar",16,0,"DWG-201",{}),
 ("P","hydrotest pressure","bar",24,0,"DWG-201",{"con":25}),
 ("P","flange thickness","mm",12.0,0.2,"DWG-201",{"dwg":12.15}),
 ("P","flange bolt circle diameter","mm",160.0,0.5,"DWG-201",{"dwg":160.4}),
 ("P","skid base length","mm",2400,5,"DWG-202",{"con":2450}),
 ("P","skid base width","mm",1200,5,"DWG-202",{"dwg":1203}),
 ("P","skid frame height","mm",850,3,"DWG-202",{}),
 ("P","anchor bolt spacing","mm",1000,2,"DWG-202",{"nocallout":True}),
 ("P","baseplate flatness","mm",0.5,0.1,"DWG-202",{}),
 ("P","piping wall thickness","mm",6.0,0.2,"DWG-301",{}),
 ("P","pipe support spacing","mm",3000,50,"DWG-301",{"dwg":3040}),
 ("P","drain connection size","mm",25,0,"DWG-301",{}),
]
# (group, parameter, contract wording, spec text, drawing note, drawing, planted-inconsistent?)
FREE = [
 ("P","gasket type","be spiral wound with graphite filler","Spiral wound, 316 stainless steel winding, graphite filler.","Gasket: ring type joint, soft iron","DWG-201",True),
 ("P","lifting arrangement","include four lifting lugs rated for the full skid weight","Four lifting lugs, rated for 1.5 times operating weight.","Lifting lugs x2, see detail","DWG-202",True),
 ("P","weld root gap","be adequate for full penetration per good industry practice","Root gap per WPS-07, 2.0 to 3.0 mm, no backing strip.","Root gap 2 mm typ., see note 4","DWG-301",False),
 ("P","surface coating","be epoxy primer with polyurethane topcoat","Epoxy primer 75 microns, polyurethane topcoat 50 microns.","Paint: epoxy primer + PU topcoat, see spec","DWG-202",False),
 ("M","nameplate marking","be stainless steel with stamped tag number","316 stainless steel nameplate, stamped tag number, riveted.","Nameplate 316SS, stamped","DWG-204",False),
 ("M","seal arrangement","be a single mechanical seal","Single mechanical seal per API 682 Plan 11.","Seal: single cartridge, Plan 11","DWG-115",False),
]
# drawing register: id, title, [(rev, date)], contract cites
DRAWINGS = {"DWG-114":("Pump Casing",[("A","2026-01-15"),("B","2026-03-02"),("C","2026-05-20")],"B","pump casing"),
 "DWG-115":("Pump Shaft and Bearing Housing",[("A","2026-02-10")],"A","pump shaft and bearing housing"),
 "DWG-201":("Flanges and Pressure Boundary",[("A","2026-02-20")],"A","flanges and pressure boundary"),
 "DWG-202":("Skid Frame",[("A","2026-02-05"),("B","2026-03-25"),("C","2026-06-01")],"B","skid frame"),
 "DWG-203":("Nozzle Arrangement",[("A","2026-02-28")],"A","nozzles"),
 "DWG-204":("Motor and Coupling",[("A","2026-03-05")],"A","motor and coupling"),
 "DWG-301":("Process Piping",[("A","2026-03-12"),("B","2026-05-08")],"B","process piping")}
CUR = {d: v[1][-1][0] for d, v in DRAWINGS.items()}
OLD = [("DWG-114","B","casing bore diameter",100.1,"mm"),("DWG-114","B","casing wall thickness",9.0,"mm"),
       ("DWG-202","B","skid base length",2400,"mm"),("DWG-301","A","piping wall thickness",5.5,"mm")]

# ---- specs (PDF) and spec ids
spec_lines = {"M": [], "P": []}; sid = {}; sec = {"M": 4, "P": 5}; cnt = {"M": 0, "P": 0}
def nextid(grp): cnt[grp] += 1; return f"{sec[grp]}.{cnt[grp]}"
for grp, p, u, sv, tol, d, o in ITEMS:
    n = nextid(grp); sid[p] = f"SPC-{grp}-{n}"
    spec_lines[grp].append(f"{n} {p.capitalize()}: {g(sv)} {u}" + (f" ±{g(tol)} {u}" if tol else "") + ".")
for grp, p, cw, st, dn, d, bad in FREE:
    n = nextid(grp); sid[p] = f"SPC-{grp}-{n}"; spec_lines[grp].append(f"{n} {p.capitalize()}: {st}")
titles = {"M": "TECHNICAL SPECIFICATION - MECHANICAL EQUIPMENT", "P": "TECHNICAL SPECIFICATION - PIPING AND STRUCTURAL"}
for grp, lines in spec_lines.items():
    c = canvas.Canvas(str(OUT / f"SPC-{grp}_{'Mechanical' if grp=='M' else 'Piping_Structural'}_Specification.pdf")); y = 800
    def put(t, bold=False):
        global y
        if y < 60: c.showPage(); y = 800
        c.setFont("Helvetica-Bold" if bold else "Helvetica", 10); c.drawString(50, y, t); y -= 15
    put(titles[grp], True); put(f"Document No: SPC-{grp}"); put("Project: Cooling Water Pump Skid Package"); put("Revision: 2"); y -= 10
    put(f"SECTION {sec[grp]} - REQUIREMENTS", True)
    for l in lines:
        wrapped = textwrap.wrap(l, 88); assert not any(re.match(r"^\d+(\.\d+)+\s", w) for w in wrapped[1:]), l
        for w in wrapped: put(w)
    c.save()

# ---- contract (docx)
doc = docx.Document(); doc.add_heading("SUPPLY CONTRACT - COOLING WATER PUMP SKID PACKAGE", 1)
for l in ["Document No: CON-001", "Employer: Eastern Process Industries Ltd.  Contractor: Meridian Fabrication Pvt. Ltd."]: doc.add_paragraph(l)
def sect(t): doc.add_paragraph(t.upper())
sect("Section 1 - General")
for l in ["1.1 The Contractor shall supply the pump skid package described in this Contract.",
          "1.2 Payment shall be made in milestones as set out in the Payment Schedule.",
          "1.3 The Contractor shall provide a warranty period of twelve months from commissioning.",
          "1.4 Liquidated damages for delay shall apply as set out in Appendix B."]: doc.add_paragraph(l)
sect("Section 6 - Technical Requirements"); k = 0; contract_num = {}
for grp, p, u, sv, tol, d, o in ITEMS:
    k += 1; contract_num[p] = f"6.{k}"
    doc.add_paragraph(f"6.{k} The {p} shall be {g(o.get('con', sv))} {u} (Ref {sid[p]}).")
for grp, p, cw, st, dn, d, bad in FREE:
    k += 1; contract_num[p] = f"6.{k}"; doc.add_paragraph(f"6.{k} The {p} shall {cw} (Ref {sid[p]}).")
sect("Section 7 - Governing Drawings")
for i, (d, (t, revs, cite, nm)) in enumerate(DRAWINGS.items(), 1):
    doc.add_paragraph(f"7.{i} Fabrication of the {nm} shall be in accordance with {d} Rev {cite}.")
doc.save(OUT / "CON-001_Supply_Contract.docx")

# ---- drawing register, callout schedule, title-block PDFs
with open(OUT / "drawing_register.csv", "w", newline="") as f:
    w = csv.writer(f); w.writerow(["drawing_id", "title", "revision", "date", "supersedes_revision"])
    for d, (t, revs, cite, nm) in DRAWINGS.items():
        for i, (r, dt) in enumerate(revs): w.writerow([d, t, r, dt, revs[i-1][0] if i else ""])
        for r, dt in revs:
            c = canvas.Canvas(str(OUT / f"{d}_Rev{r}.pdf")); c.rect(30, 30, 535, 780); c.rect(330, 30, 235, 90)
            c.setFont("Helvetica-Bold", 11); c.drawString(340, 100, f"{d}  Rev {r}"); c.setFont("Helvetica", 9)
            c.drawString(340, 85, t); c.drawString(340, 70, f"Date: {dt}"); c.drawString(340, 55, "Content not extracted (stub)"); c.save()
rows = []; n = 0
def cal(d, r, p, v, u, note): 
    global n; n += 1; rows.append([f"CAL-{n:03d}", d, r, p, "" if v is None else g(v), u, note])
for grp, p, u, sv, tol, d, o in ITEMS:
    if not o.get("nocallout"): cal(d, CUR[d], p, o.get("dwg", sv), u, f"{p} {g(o.get('dwg', sv))} {u}")
for d, r, p, v, u in OLD: cal(d, r, p, v, u, f"{p} {g(v)} {u}")
for grp, p, cw, st, dn, d, bad in FREE: cal(d, CUR[d], p, None, "", dn)
with open(OUT / "callout_schedule.csv", "w", newline="") as f:
    w = csv.writer(f); w.writerow(["callout_id", "drawing_id", "revision", "parameter", "value", "unit", "note"]); w.writerows(rows)

fails = [p for grp, p, u, sv, tol, d, o in ITEMS if "con" in o or ("dwg" in o and abs(o["dwg"] - sv) > tol + 1e-9)]
json.dump({"expected_numeric_fail": fails,
           "expected_revision_fail": [f"governing revision {d.lower()}" for d, v in DRAWINGS.items() if v[2] != CUR[d]],
           "expected_semantic_inconsistent": [p for _, p, *_r, bad in FREE if bad],
           "expected_coverage_gaps": [p for grp, p, u, sv, tol, d, o in ITEMS if o.get("nocallout")],
           "expected_superseded_excluded": len(OLD)}, open(OUT / "_manifest.json", "w"), indent=1)
print(len(list(OUT.iterdir())), "files;", n, "callouts;", fails)
