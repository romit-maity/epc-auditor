"""DocumentSource over a folder of real files: contracts/specs (.docx .pdf .txt), drawing register + callout schedule (.csv).

Clause grammar (deterministic regex, deliberately strict; anything that doesn't match is counted and reported, never guessed):
  contract  "6.1 The <parameter> shall be <number> <unit> (Ref SPC-M-4.1)."   numeric
            "6.9 The <parameter> shall <free text> (Ref SPC-P-5.3)."          free text
            "7.1 ... in accordance with DWG-114 Rev B."                        revision citation
  spec      "4.1 <Parameter>: <number> <unit> [±<tol> <unit>]."  |  "5.3 <Parameter>: <free text>"
"""
import csv, re
from pathlib import Path
from models import Corpus, Document, Clause, SpecReq, DrawingRev, Callout
from ingestion.source_interface import DocumentSource
from ingestion.drawing_extraction import extract_callouts

NUM = r"-?\d+(?:\.\d+)?"
R_CLAUSE = re.compile(r"^(?P<n>\d+(?:\.\d+)+)\s+(?P<body>.+)$")
R_DOCNO = re.compile(r"^Document No:\s*(\S+)")
R_REF = re.compile(r"\(Ref (?P<ref>[A-Z]+-[A-Z0-9]+-\d+(?:\.\d+)+)\)")
R_REV = re.compile(r"in accordance with (?P<d>DWG-\d+) Rev (?P<r>[A-Z])")
R_CNUM = re.compile(rf"^The (?P<p>.+?) shall be (?P<v>{NUM})\s?(?P<u>[A-Za-z]+)\b")
R_CFREE = re.compile(r"^The (?P<p>.+?) shall (?P<t>.+?)\s*\(Ref")
R_SNUM = re.compile(rf"^(?P<p>[^:]+):\s*(?P<v>{NUM})\s*(?P<u>[A-Za-z]+)(?:\s*±\s*(?P<t>{NUM})\s*[A-Za-z]*)?\s*\.?$")
R_SFREE = re.compile(r"^(?P<p>[^:]+):\s*(?P<t>.+)$")

def read_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".docx":
        import docx; return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
    if ext == ".pdf":
        from pypdf import PdfReader; return "\n".join(pg.extract_text() or "" for pg in PdfReader(str(path)).pages)
    return path.read_text(errors="ignore")

def split_clauses(text: str):
    """Return (title, doc_no, [(number, full clause text)]); wrapped lines are joined to their clause."""
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    title, doc_no, out = lines[0] if lines else "", None, []
    for l in lines[1:]:
        if (m := R_DOCNO.match(l)): doc_no = m.group(1)
        elif (m := R_CLAUSE.match(l)): out.append([m.group("n"), m.group("body")])
        elif out and not l.isupper(): out[-1][1] += " " + l
    return title, doc_no, out

class FileSource(DocumentSource):
    def __init__(self, folder: str): self.folder = Path(folder)

    def load(self) -> Corpus:
        c = dict(documents=[], clauses=[], spec_reqs=[], drawing_revs=[], callouts=[]); warn = []
        st = dict(files=0, contracts=0, specs=0, clauses_parsed=0, clauses_ignored=0, drawing_sheets_stubbed=0)
        files = sorted(p for p in self.folder.iterdir() if p.is_file() and not p.name.startswith("_"))
        for p in files:
            st["files"] += 1
            if p.suffix.lower() == ".csv": self._csv(p, c, warn)
            elif p.name.upper().startswith("DWG-"): extract_callouts(p); st["drawing_sheets_stubbed"] += 1
            elif p.suffix.lower() in (".docx", ".pdf", ".txt", ".md"): self._doc(p, c, warn, st)
            else: warn.append(f"{p.name}: unsupported file type, skipped")
        ids = {s.id for s in c["spec_reqs"]}; revs = {r.id for r in c["drawing_revs"]}
        for x in c["clauses"]:
            if x.spec_ref and x.spec_ref not in ids: warn.append(f"{x.id}: references unknown spec {x.spec_ref}")
            if x.cites_drawing_rev and x.cites_drawing_rev not in revs: warn.append(f"{x.id}: cites {x.cites_drawing_rev}, not in drawing register")
        for k in c["callouts"]:
            if k.drawing_rev_id not in revs: warn.append(f"{k.id}: callout on unregistered revision {k.drawing_rev_id}")
        return Corpus(**c, warnings=warn, stats=st)

    def _doc(self, p, c, warn, st):
        title, doc_no, cl = split_clauses(read_text(p))
        kind = "contract" if "CONTRACT" in title.upper() else "spec" if "SPECIFICATION" in title.upper() else None
        if not kind or not doc_no: warn.append(f"{p.name}: not recognised as contract/spec (need CONTRACT/SPECIFICATION title and 'Document No:')"); return
        c["documents"].append(Document(id=doc_no, kind=kind, title=title)); st["contracts" if kind == "contract" else "specs"] += 1
        for n, body in cl:
            if kind == "spec":
                if (m := R_SNUM.match(body)):
                    c["spec_reqs"].append(SpecReq(id=f"{doc_no}-{n}", doc_id=doc_no, text=f"{n} {body}", parameter=m["p"].strip(), kind="numeric",
                                                  value=float(m["v"]), tol=float(m["t"] or 0), unit=m["u"]))
                elif (m := R_SFREE.match(body)):
                    c["spec_reqs"].append(SpecReq(id=f"{doc_no}-{n}", doc_id=doc_no, text=f"{n} {body}", parameter=m["p"].strip(), kind="free_text"))
                else: st["clauses_ignored"] += 1; continue
            else:
                cid, ref = f"{doc_no}:{n}", R_REF.search(body)
                if (m := R_REV.search(body)):
                    c["clauses"].append(Clause(id=cid, doc_id=doc_no, text=f"{n} {body}", parameter=f"Governing revision {m['d']}", kind="revision",
                                               cites_drawing_rev=f"{m['d']}@{m['r']}"))
                elif ref and (m := R_CNUM.match(body)):
                    c["clauses"].append(Clause(id=cid, doc_id=doc_no, text=f"{n} {body}", parameter=m["p"], kind="numeric", value=float(m["v"]),
                                               unit=m["u"], spec_ref=ref["ref"]))
                elif ref and (m := R_CFREE.match(body)):
                    c["clauses"].append(Clause(id=cid, doc_id=doc_no, text=f"{n} {body}", parameter=m["p"], kind="free_text", spec_ref=ref["ref"]))
                else: st["clauses_ignored"] += 1; continue
            st["clauses_parsed"] += 1

    def _csv(self, p, c, warn):
        rows = list(csv.DictReader(p.open(newline="", encoding="utf-8-sig")))
        cols = set(rows[0]) if rows else set()
        if {"drawing_id", "title", "revision", "date", "supersedes_revision"} <= cols:
            seen = set()
            for r in rows:
                if r["drawing_id"] not in seen:
                    seen.add(r["drawing_id"]); c["documents"].append(Document(id=r["drawing_id"], kind="drawing", title=r["title"]))
                sup = r["supersedes_revision"].strip()
                c["drawing_revs"].append(DrawingRev(id=f"{r['drawing_id']}@{r['revision']}", drawing_id=r["drawing_id"], revision=r["revision"],
                                                    date=r["date"], supersedes=f"{r['drawing_id']}@{sup}" if sup else None))
        elif {"callout_id", "drawing_id", "revision", "parameter", "value", "unit", "note"} <= cols:
            for r in rows:
                c["callouts"].append(Callout(id=r["callout_id"], drawing_rev_id=f"{r['drawing_id']}@{r['revision']}", parameter=r["parameter"],
                                             value=float(r["value"]) if r["value"].strip() else None, unit=r["unit"] or None, text=r["note"]))
        else: warn.append(f"{p.name}: CSV headers match neither drawing register nor callout schedule, skipped")
